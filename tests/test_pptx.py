"""Needs: pip install python-pptx pillow. Run: python -m unittest discover -s tests"""
import importlib.util
import subprocess
import sys
import tempfile
import unittest
from pathlib import Path

from PIL import Image
from pptx import Presentation
from pptx.util import Inches, Pt

SCRIPTS = Path(__file__).resolve().parents[1] / "skills" / "zero1-pptx" / "scripts"
# each skill has its own helpers.py: load by path under a unique name so the three test modules do not clash
_spec = importlib.util.spec_from_file_location("zero1_pptx_helpers", SCRIPTS / "helpers.py")
h = importlib.util.module_from_spec(_spec)
_spec.loader.exec_module(h)


def run(script, *args):
    return subprocess.run([sys.executable, str(SCRIPTS / script), *map(str, args)], capture_output=True, text=True)


class PptxScripts(unittest.TestCase):
    def setUp(self):
        self.tmp = Path(tempfile.mkdtemp())
        self.png = self.tmp / "p.png"
        Image.new("RGB", (400, 200), "orange").save(self.png)

    def build(self):
        prs = h.new_deck()
        title = prs.slides.add_slide(prs.slide_layouts[0])
        title.shapes.title.text = "Report {{periodo}}"
        title.placeholders[1].text = "Per {{cliente}}"
        s = prs.slides.add_slide(prs.slide_layouts[1])
        h.set_title(s, "Punti chiave")
        h.add_bullets(s, ["Primo", ("Dettaglio", 1), "Secondo"], size_pt=20)
        h.set_notes(s, "Parlare dei ricavi.")
        c = prs.slides.add_slide(prs.slide_layouts[5])
        h.set_title(c, "Ricavi")
        h.add_bar_chart(c, ["Q1", "Q2"], {"2026": [10, 12]}, 0.8, 1.6, 6, 4.5, title="Ricavi (M€)")
        h.add_picture_fit(c, self.png, 7.2, 1.6, 5, 4.5, "Logo arancione")
        t = prs.slides.add_slide(prs.slide_layouts[5])
        h.set_title(t, "Tabella")
        h.add_table(t, ["Voce", "Importo"], [["A", "10"], ["B", "20"]], 0.8, 1.8, 8, align_right=[1])
        return prs

    def test_helpers_make_a_deck_that_passes_the_check(self):
        prs = self.build()
        path = self.tmp / "deck.pptx"
        prs.save(path)
        again = Presentation(path)
        self.assertEqual(round(again.slide_width.inches, 2), 13.33)
        # the layouts followed the widescreen resize: no placeholder beyond the page
        for layout in again.slide_layouts:
            for shape in layout.shapes:
                if shape.left is not None:
                    self.assertLessEqual(shape.left + shape.width, again.slide_width)
        r = run("describe.py", path, "--check")
        self.assertEqual(r.returncode, 0, r.stdout + r.stderr)
        self.assertIn("notes=yes", r.stdout)

    def test_check_flags_unfinished_slides(self):
        prs = h.new_deck()
        s = prs.slides.add_slide(prs.slide_layouts[1])      # title and body left empty
        box = s.shapes.add_textbox(Inches(11), Inches(1), Inches(4), Inches(0.5))   # beyond the right edge
        box.text_frame.text = "x"
        small = s.shapes.add_textbox(Inches(1), Inches(5), Inches(3), Inches(0.5))
        small.text_frame.text = "piccolo"
        small.text_frame.paragraphs[0].runs[0].font.size = Pt(8)
        s.shapes.add_picture(str(self.png), Inches(1), Inches(1), Inches(2))
        long = s.shapes.add_textbox(Inches(1), Inches(6), Inches(2), Inches(0.4))
        long.text_frame.word_wrap = True
        long.text_frame.text = "parola " * 130
        path = self.tmp / "bad.pptx"
        prs.save(path)
        r = run("describe.py", path, "--check")
        self.assertEqual(r.returncode, 1)
        for needle in ("no title", "empty placeholder", "beyond the slide", "smaller than 12", "no alt text", "words", "overflows"):
            self.assertIn(needle, r.stdout)

    def test_replace_text_across_runs_in_tables_and_notes(self):
        prs = self.build()
        p = prs.slides[0].placeholders[1].text_frame.paragraphs[0]
        p.runs[0].text = "Per {{cli"            # split the marker over two runs with different formatting
        extra = p.add_run()
        extra.text = "ente}}"
        extra.font.bold = True
        h.set_notes(prs.slides[1], "Cliente: {{cliente}} A")
        src, dst = self.tmp / "t.pptx", self.tmp / "o.pptx"
        prs.save(src)
        r = run("replace_text.py", src, dst, "--set", "{{cliente}}=ACME", "--set", "{{periodo}}=Q3", "--set", "A=AA")
        self.assertEqual(r.returncode, 0, r.stdout + r.stderr)
        out = Presentation(dst)
        self.assertEqual(out.slides[0].shapes.title.text, "Report Q3")
        # "A" -> "AA" runs after {{cliente}} -> ACME and must neither loop nor rescan its own output
        self.assertEqual(out.slides[0].placeholders[1].text_frame.text, "Per AACME")
        self.assertEqual(out.slides[1].notes_slide.notes_text_frame.text, "Cliente: AACME AA")
        self.assertEqual(run("replace_text.py", src, dst, "--set", "{{assente}}=x").returncode, 1)

    def test_duplicate_move_delete_keep_pictures_working(self):
        prs = self.build()
        before = len(prs.slides)
        h.duplicate_slide(prs, 2)             # slide with chart and picture
        self.assertEqual(len(prs.slides), before + 1)
        h.move_slide(prs, before, 1)
        h.delete_slide(prs, 0)
        path = self.tmp / "dup.pptx"
        prs.save(path)
        again = Presentation(path)
        self.assertEqual(len(again.slides), before)
        pics = [sh for sl in again.slides for sh in sl.shapes if sh.shape_type == 13]
        self.assertEqual(len(pics), 2)
        for pic in pics:
            self.assertGreater(len(pic.image.blob), 100)   # the copy points at a real image part
        r = run("describe.py", path)
        self.assertEqual(r.returncode, 0, r.stderr)

    def test_estimate_font_size_shrinks_with_text(self):
        self.assertGreater(h.estimate_font_size("Breve", 6, 1), h.estimate_font_size("lungo " * 80, 6, 1))
        self.assertEqual(h.estimate_font_size("x " * 2000, 2, 0.5, min_pt=12), 12)

    def test_render_without_libreoffice_fails_with_a_clear_message(self):
        r = run("render.py", self.tmp / "missing.pptx")
        if r.returncode == 0:
            self.skipTest("LibreOffice present")
        self.assertIn("LibreOffice", r.stdout + r.stderr)


if __name__ == "__main__":
    unittest.main()
