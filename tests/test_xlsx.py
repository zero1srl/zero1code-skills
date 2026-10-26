"""Needs: pip install openpyxl. Run: python -m unittest discover -s tests"""
import importlib.util
import subprocess
import sys
import tempfile
import unittest
from pathlib import Path

import openpyxl
from openpyxl.chart import BarChart, Reference
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

SCRIPTS = Path(__file__).resolve().parents[1] / "skills" / "zero1-xlsx" / "scripts"
# each skill has its own helpers.py: load by path under a unique name so the three test modules do not clash
_spec = importlib.util.spec_from_file_location("zero1_xlsx_helpers", SCRIPTS / "helpers.py")
h = importlib.util.module_from_spec(_spec)
_spec.loader.exec_module(h)


def run(script, *args):
    return subprocess.run([sys.executable, str(SCRIPTS / script), *map(str, args)], capture_output=True, text=True)


class XlsxScripts(unittest.TestCase):
    def setUp(self):
        self.tmp = Path(tempfile.mkdtemp())

    def test_helpers_build_a_workbook_that_passes_the_check(self):
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Dati"
        h.write_rows(ws, [["Voce", "Quantità", "Prezzo", "Totale"], ["A", 2, 10.5, "=B2*C2"], ["B", 3, 7, "=B3*C3"],
                          ["Totale", "=SUM(B2:B3)", "", "=SUM(D2:D3)"]])
        h.style_header(ws)
        for r in range(2, 5):
            ws.cell(row=r, column=3).number_format = h.EUR
            ws.cell(row=r, column=4).number_format = h.EUR
        h.mark_input(ws["B2"], ws["B3"], ws["C2"], ws["C3"])
        h.autofit_columns(ws)
        h.fit_to_page(ws)
        dv = DataValidation(type="list", formula1='"A,B"', allow_blank=True)
        ws.add_data_validation(dv)
        dv.add("A2:A3")
        ws.conditional_formatting.add("D2:D3", CellIsRule(operator="greaterThan", formula=["20"], fill=PatternFill("solid", bgColor="FFC7CE")))
        chart = BarChart()
        chart.add_data(Reference(ws, min_col=4, min_row=1, max_row=3), titles_from_data=True)
        ws.add_chart(chart, "F2")
        summary = wb.create_sheet("Riepilogo")
        summary["A1"] = "Totale"
        summary["B1"] = "=Dati!D4"
        from openpyxl.workbook.defined_name import DefinedName
        summary["A2"] = "Aliquota"
        summary["B2"] = 0.22
        wb.defined_names["Aliquota"] = DefinedName("Aliquota", attr_text="Riepilogo!$B$2")
        wb.calculation.fullCalcOnLoad = True
        path = self.tmp / "ok.xlsx"
        wb.save(path)
        r = run("describe.py", path, "--check")
        self.assertEqual(r.returncode, 0, r.stdout + r.stderr)
        self.assertIn("defined names: Aliquota", r.stdout)
        self.assertIn("freeze=A2", r.stdout)
        self.assertIn("charts=1", r.stdout)
        self.assertIn("no cached results yet", r.stdout)
        self.assertIn("contains 1 charts", r.stdout)  # a chart written by openpyxl is a part openpyxl cannot read back

    def test_check_flags_errors_missing_sheets_and_text_numbers(self):
        wb = openpyxl.Workbook()
        ws = wb.active
        ws["A1"] = "Valore"
        ws["A2"] = "1.234,50"          # number stored as text
        ws["B2"] = "=Mancante!A1"      # sheet does not exist
        ws["B3"] = "=#REF!+1"
        ws["B4"] = '="Fatto!"'          # a '!' inside a string is not a sheet reference
        ws["C1"] = "#DIV/0!"           # error value stored in the cell
        path = self.tmp / "bad.xlsx"
        wb.save(path)
        r = run("describe.py", path, "--check")
        self.assertEqual(r.returncode, 1)
        for needle in ("text that looks like a number", "does not exist: 'Mancante'", "contains #REF!", "#DIV/0!"):
            self.assertIn(needle, r.stdout)
        self.assertNotIn("Fatto", r.stdout.split("== check")[1])

    def test_recalc_without_libreoffice_says_so(self):
        r = run("recalc.py", self.tmp / "x.xlsx")
        if r.returncode == 0:
            self.skipTest("LibreOffice present")
        self.assertIn("LibreOffice", r.stdout + r.stderr)


if __name__ == "__main__":
    unittest.main()
