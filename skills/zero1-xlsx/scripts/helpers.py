"""Small openpyxl helpers for workbooks that look finished.

Use from a script:  sys.path.insert(0, "<this folder>"); from helpers import *
"""
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

__all__ = ["EUR", "EUR0", "PCT", "DATE", "INT", "write_rows", "style_header", "autofit_columns", "add_table",
           "mark_input", "fit_to_page", "thin_borders"]

# Number formats are stored in US notation; Excel shows them with the reader's own separators.
EUR = '#,##0.00 "€"'
EUR0 = '#,##0 "€"'
PCT = "0.0%"
DATE = "dd/mm/yyyy"
INT = "#,##0"

INPUT_FONT = Font(color="0000FF")  # blue = typed-in assumption, black = formula (financial-model convention)


def write_rows(ws, rows, start_row=1, start_col=1):
    """Write a list of lists; strings starting with '=' become formulas. Returns the last row written."""
    r = start_row
    for r, values in enumerate(rows, start=start_row):
        for c, value in enumerate(values, start=start_col):
            ws.cell(row=r, column=c, value=value)
    return r


def style_header(ws, row=1, fill="1F4E78", font_color="FFFFFF", freeze=True):
    """Bold coloured header over the used columns; optionally freezes the rows above and including it."""
    for cell in ws[row]:
        if cell.value is not None:
            cell.font = Font(bold=True, color=font_color)
            cell.fill = PatternFill("solid", fgColor=fill)
            cell.alignment = Alignment(vertical="center", wrap_text=True)
    if freeze:
        ws.freeze_panes = ws.cell(row=row + 1, column=1)


def autofit_columns(ws, min_width=8, max_width=60):
    """Column width from the longest displayed value; formulas count as 12 characters."""
    widths = {}
    for row in ws.iter_rows():
        for cell in row:
            if cell.value is None:
                continue
            text = "x" * 12 if cell.data_type == "f" else str(cell.value)
            longest = max(len(line) for line in text.splitlines() or [""])
            widths[cell.column_letter] = max(widths.get(cell.column_letter, 0), longest)
    for letter, size in widths.items():
        ws.column_dimensions[letter].width = max(min_width, min(max_width, size + 2))


def add_table(ws, first_cell="A1", name="Dati", style="TableStyleMedium2"):
    """Turn the block that starts at `first_cell` into an Excel table (filters, banding, structured references)."""
    start = ws[first_cell]
    ref = f"{start.coordinate}:{get_column_letter(ws.max_column)}{ws.max_row}"
    table = Table(displayName=name, ref=ref)
    table.tableStyleInfo = TableStyleInfo(name=style, showRowStripes=True)
    ws.add_table(table)
    return table


def mark_input(*cells):
    """Blue font for cells the user is meant to change."""
    for cell in cells:
        cell.font = INPUT_FONT


def thin_borders(ws, ref):
    side = Side(style="thin", color="BFBFBF")
    for row in ws[ref]:
        for cell in row:
            cell.border = Border(left=side, right=side, top=side, bottom=side)


def fit_to_page(ws, landscape=True):
    """Print on A4, one page wide, header row repeated."""
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.orientation = "landscape" if landscape else "portrait"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_title_rows = "1:1"
