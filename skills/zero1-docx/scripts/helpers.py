"""python-docx has no API for these common needs; each helper adds the missing WordprocessingML.

Use from a script:  sys.path.insert(0, "<this folder>"); from helpers import *
Needs python-docx >= 1.2 (comments come from the library itself: doc.add_comment).
"""
from docx.enum.section import WD_ORIENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Mm, Pt, RGBColor

__all__ = [
    "a4", "landscape", "set_base_font", "set_language", "add_field", "add_page_number_footer", "add_toc",
    "update_fields_on_open", "add_hyperlink", "shade_cell", "repeat_header", "add_table", "add_image", "page_break",
]


def a4(doc, margin_mm=25):
    """A4 portrait with equal margins on every section (python-docx starts from US Letter)."""
    for s in doc.sections:
        s.orientation = WD_ORIENT.PORTRAIT
        s.page_width, s.page_height = Mm(210), Mm(297)
        s.left_margin = s.right_margin = s.top_margin = s.bottom_margin = Mm(margin_mm)


def landscape(section):
    """Turn one section to landscape: swap the page size as well as the orientation flag."""
    section.orientation = WD_ORIENT.LANDSCAPE
    w, h = section.page_width, section.page_height
    if w < h:
        section.page_width, section.page_height = h, w


def set_base_font(doc, name="Calibri", size_pt=11):
    """Body font for the Normal style, including the East Asian and complex-script slots Word reads."""
    style = doc.styles["Normal"]
    style.font.name = name
    style.font.size = Pt(size_pt)
    fonts = style.element.get_or_add_rPr().get_or_add_rFonts()
    for slot in ("w:eastAsia", "w:cs"):
        fonts.set(qn(slot), name)


def set_language(doc, tag="it-IT"):
    """Proofing language for the whole document (spell check and hyphenation follow it)."""
    defaults = doc.styles.element.xpath("w:docDefaults/w:rPrDefault/w:rPr")
    rpr = defaults[0] if defaults else doc.styles.element.get_or_add_docDefaults().get_or_add_rPrDefault().get_or_add_rPr()
    lang = rpr.find(qn("w:lang"))
    if lang is None:
        lang = OxmlElement("w:lang")
        rpr.append(lang)
    for slot in ("w:val", "w:eastAsia", "w:bidi"):
        lang.set(qn(slot), tag)


def add_field(paragraph, instruction, placeholder="1"):
    """Insert a Word field (PAGE, NUMPAGES, DATE, TOC ...). `placeholder` is what shows until Word updates it."""
    def mark(kind):
        run = paragraph.add_run()
        char = OxmlElement("w:fldChar")
        char.set(qn("w:fldCharType"), kind)
        run._r.append(char)

    mark("begin")
    run = paragraph.add_run()
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = f" {instruction} "
    run._r.append(instr)
    mark("separate")
    paragraph.add_run(placeholder)
    mark("end")


def add_page_number_footer(section, prefix="Pagina "):
    """Centered 'Pagina X di Y' in the section footer (unlinks it from the previous section's footer)."""
    footer = section.footer
    footer.is_linked_to_previous = False
    p = footer.paragraphs[0]
    p.text = ""
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run(prefix)
    add_field(p, "PAGE")
    p.add_run(" di ")
    add_field(p, "NUMPAGES")


def add_toc(doc, levels="1-3", hint="Aggiorna il campo (clic destro, Aggiorna campo) per generare l'indice."):
    """Table of contents field built from Heading styles. Word fills it in when fields are updated."""
    p = doc.add_paragraph()
    add_field(p, f'TOC \\o "{levels}" \\h \\z \\u', placeholder=hint)
    return p


def update_fields_on_open(doc):
    """Ask Word to refresh fields (TOC, page counts) when the file is opened; Word shows a one-time prompt."""
    settings = doc.settings.element
    if settings.find(qn("w:updateFields")) is None:
        el = OxmlElement("w:updateFields")
        el.set(qn("w:val"), "true")
        settings.append(el)


def add_hyperlink(paragraph, url, text):
    """External link with the usual blue underline (the default template has no Hyperlink character style)."""
    rid = paragraph.part.relate_to(
        url, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink", is_external=True)
    link = OxmlElement("w:hyperlink")
    link.set(qn("r:id"), rid)
    run = OxmlElement("w:r")
    rpr = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "0563C1")
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    rpr.extend([color, underline])
    run.append(rpr)
    t = OxmlElement("w:t")
    t.text = text
    t.set(qn("xml:space"), "preserve")
    run.append(t)
    link.append(run)
    paragraph._p.append(link)
    return link


def shade_cell(cell, fill_hex):
    """Background colour of a table cell, e.g. 'D9E2F3'."""
    tc_pr = cell._tc.get_or_add_tcPr()
    for old in tc_pr.findall(qn("w:shd")):
        tc_pr.remove(old)
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill_hex)
    tc_pr.append(shd)


def repeat_header(row):
    """Mark a table row as the header that repeats on every page."""
    tr_pr = row._tr.get_or_add_trPr()
    if tr_pr.find(qn("w:tblHeader")) is None:
        el = OxmlElement("w:tblHeader")
        el.set(qn("w:val"), "true")
        tr_pr.append(el)


def add_table(doc, header, rows, widths_mm=None, style="Table Grid", header_fill="D9E2F3", align_right=()):
    """Table with a bold shaded repeating header. `align_right` lists the column indexes to right-align (numbers)."""
    table = doc.add_table(rows=1, cols=len(header))
    table.style = style
    table.autofit = False
    for i, text in enumerate(header):
        cell = table.rows[0].cells[i]
        cell.text = ""
        cell.paragraphs[0].add_run(str(text)).bold = True
        shade_cell(cell, header_fill)
    repeat_header(table.rows[0])
    for values in rows:
        cells = table.add_row().cells
        for i, value in enumerate(values):
            cells[i].text = str(value)
    for row in table.rows:
        for i in align_right:
            row.cells[i].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.RIGHT
        if widths_mm:
            for cell, width in zip(row.cells, widths_mm):
                cell.width = Mm(width)
    return table


def add_image(doc, path, width_mm, alt, align=WD_ALIGN_PARAGRAPH.CENTER):
    """Picture in its own paragraph, with alt text (screen readers and the --check in describe.py need it)."""
    p = doc.add_paragraph()
    p.alignment = align
    shape = p.add_run().add_picture(str(path), width=Mm(width_mm))
    shape._inline.docPr.set("descr", alt)
    return p


def page_break(doc):
    """Hard page break in its own paragraph; prefer 'page break before' on a heading style when possible."""
    from docx.enum.text import WD_BREAK
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)


def color(run, hex_rgb):
    """Run font colour from 'RRGGBB'."""
    run.font.color.rgb = RGBColor.from_string(hex_rgb.upper())
