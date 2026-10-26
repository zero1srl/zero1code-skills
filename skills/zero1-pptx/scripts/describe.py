#!/usr/bin/env python3
"""Describe a .pptx and, with --check, list problems that make a deck look unfinished.

  python describe.py FILE.pptx [--check]

Per slide: layout, title, every shape with position (inches), type and text excerpt, notes. --check exits with 1 on
shapes outside the slide, empty placeholders, slides without a title, text that probably overflows its box,
text smaller than 12 pt, pictures without alt text, slides with more than 120 words, and repeated titles.
Overflow is an estimate (average glyph width 0.5 em): confirm with render.py.
"""
import argparse
import math
import re
import sys
from collections import Counter

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE, PP_PLACEHOLDER
from pptx.util import Emu

MIN_PT = 12
IMAGE_NAME = re.compile(r".(png|jpe?g|gif|bmp|tiff?|svg|emf|wmf|webp)$", re.I)  # python-pptx fills alt text with the file name


def has_alt(descr):
    return bool(descr.strip()) and not IMAGE_NAME.search(descr.strip())
MAX_WORDS = 120


def inches(value):
    return Emu(value).inches if value is not None else 0.0


def walk(shapes):
    for shape in shapes:
        yield shape
        if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
            yield from walk(shape.shapes)


def text_of(shape):
    return shape.text_frame.text if getattr(shape, "has_text_frame", False) and shape.has_text_frame else ""


def estimated_overflow(shape):
    """True when the text, at its largest explicit size (or 18 pt), needs more height than the box has."""
    frame = shape.text_frame
    if not frame.text.strip() or shape.width is None or shape.height is None:
        return False
    sizes = [r.font.size.pt for p in frame.paragraphs for r in p.runs if r.font.size is not None]
    size = max(sizes) if sizes else 18
    width_pt = inches(shape.width) * 72 - 14.4      # default left/right inset 0.1 in each
    chars = max(1, int(width_pt / (size * 0.5)))
    lines = sum(max(1, math.ceil(len(p.text) / chars)) for p in frame.paragraphs)
    return lines * size * 1.2 > inches(shape.height) * 72 - 7.2


def describe(prs):
    print(f"== deck: {len(prs.slides)} slide(s), {inches(prs.slide_width):.2f} x {inches(prs.slide_height):.2f} in")
    for number, slide in enumerate(prs.slides, start=1):
        title = slide.shapes.title.text if slide.shapes.title is not None else None
        notes = slide.has_notes_slide and slide.notes_slide.notes_text_frame.text.strip()
        print(f"-- slide {number} [{slide.slide_layout.name}] title={title!r} notes={'yes' if notes else 'no'}")
        for shape in walk(slide.shapes):
            kind = str(shape.shape_type).split(".")[-1].split(" ")[0] if shape.shape_type else "?"
            excerpt = text_of(shape).replace("\n", " / ")[:60]
            print(f"   {shape.name[:24]:24} {kind:14} ({inches(shape.left):.1f},{inches(shape.top):.1f}) "
                  f"{inches(shape.width):.1f}x{inches(shape.height):.1f} {excerpt}")


def check(prs):
    problems = []
    width, height = prs.slide_width, prs.slide_height
    titles = Counter()
    for number, slide in enumerate(prs.slides, start=1):
        at = f"slide {number}"
        title = slide.shapes.title.text.strip() if slide.shapes.title is not None else ""
        if not title:
            problems.append(f"{at}: no title (needed for navigation and accessibility)")
        else:
            titles[title] += 1
        words = 0
        for shape in walk(slide.shapes):
            if shape.left is not None and shape.width is not None:
                if shape.left < 0 or shape.top < 0 or shape.left + shape.width > width or shape.top + shape.height > height:
                    problems.append(f"{at}: '{shape.name}' extends beyond the slide")
            if shape.is_placeholder and getattr(shape, "has_text_frame", False) and shape.has_text_frame \
                    and not shape.text_frame.text.strip() \
                    and shape.placeholder_format.type not in (PP_PLACEHOLDER.SLIDE_NUMBER, PP_PLACEHOLDER.DATE, PP_PLACEHOLDER.FOOTER):
                problems.append(f"{at}: empty placeholder '{shape.name}' (fill it or delete it)")
            if getattr(shape, "has_text_frame", False) and shape.has_text_frame:
                words += len(shape.text_frame.text.split())
                if estimated_overflow(shape):
                    problems.append(f"{at}: text in '{shape.name}' probably overflows its box")
                small = {r.font.size.pt for p in shape.text_frame.paragraphs for r in p.runs
                         if r.font.size is not None and r.font.size.pt < MIN_PT and r.text.strip()}
                if small:
                    problems.append(f"{at}: '{shape.name}' uses text smaller than {MIN_PT} pt ({sorted(small)})")
            if shape.shape_type == MSO_SHAPE_TYPE.PICTURE and not has_alt(shape._element.nvPicPr.cNvPr.get("descr") or ""):
                problems.append(f"{at}: picture '{shape.name}' has no alt text")
        if words > MAX_WORDS:
            problems.append(f"{at}: {words} words (over {MAX_WORDS}); split the slide or move detail to the notes")
    for text, count in titles.items():
        if count > 1:
            problems.append(f"title repeated on {count} slides: {text!r}")
    return problems


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("file")
    ap.add_argument("--check", action="store_true")
    args = ap.parse_args()
    prs = Presentation(args.file)
    describe(prs)
    if args.check:
        problems = check(prs)
        print(f"== check: {len(problems)} problem(s)")
        for line in problems:
            print("  -", line)
        sys.exit(1 if problems else 0)


if __name__ == "__main__":
    main()
