#!/usr/bin/env python3
"""Fill a template deck: replace text everywhere (slides, tables, groups, notes) keeping the formatting.

  python replace_text.py IN.pptx OUT.pptx --set "{{cliente}}=ACME" --set "{{data}}=26/10/2026"
  python replace_text.py IN.pptx OUT.pptx --map values.json          # {"{{cliente}}": "ACME", ...}

A marker split across runs is handled: the replacement takes the formatting of the run where the marker starts.
Prints each marker with the number of replacements and exits 1 if any marker was not found.
"""
import argparse
import json
import sys
from collections import Counter

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE


def paragraphs(shapes):
    for shape in shapes:
        if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
            yield from paragraphs(shape.shapes)
        if getattr(shape, "has_text_frame", False) and shape.has_text_frame:
            yield from shape.text_frame.paragraphs
        if getattr(shape, "has_table", False) and shape.has_table:
            for row in shape.table.rows:
                for cell in row.cells:
                    yield from cell.text_frame.paragraphs


def replace_in_paragraph(paragraph, old, new):
    count, search_from = 0, 0
    while True:
        runs = list(paragraph.runs)
        text = "".join(r.text for r in runs)
        start = text.find(old, search_from)
        if start < 0:
            return count
        end = start + len(old)
        pos, first, last = 0, None, None
        for i, run in enumerate(runs):
            length = len(run.text)
            if first is None and start < pos + length:
                first, first_off = i, start - pos
            if first is not None and end <= pos + length:
                last, last_off = i, end - pos
                break
            pos += length
        head, tail = runs[first].text[:first_off], runs[last].text[last_off:]
        if first == last:
            runs[first].text = head + new + tail
        else:
            runs[first].text = head + new
            runs[last].text = tail
            for run in runs[first + 1:last]:
                run._r.getparent().remove(run._r)
        search_from = start + len(new)  # never re-scan the text just inserted
        count += 1


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("src")
    ap.add_argument("dst")
    ap.add_argument("--set", action="append", default=[], metavar="OLD=NEW")
    ap.add_argument("--map", help="JSON file with an object of OLD: NEW")
    args = ap.parse_args()
    values = json.load(open(args.map, encoding="utf-8")) if args.map else {}
    for item in args.set:
        old, _, new = item.partition("=")
        values[old] = new
    if not values:
        ap.error("give --set or --map")
    prs = Presentation(args.src)
    found = Counter()
    for slide in prs.slides:
        scopes = [slide.shapes]
        if slide.has_notes_slide:
            scopes.append(slide.notes_slide.shapes)
        for scope in scopes:
            for paragraph in paragraphs(scope):
                for old, new in values.items():
                    found[old] += replace_in_paragraph(paragraph, old, new)
    prs.save(args.dst)
    missing = [k for k in values if not found[k]]
    for old in values:
        print(f"{found[old]:3} x {old}")
    sys.exit(1 if missing else 0)


if __name__ == "__main__":
    main()
