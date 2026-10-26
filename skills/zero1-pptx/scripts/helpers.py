"""python-pptx helpers for decks that look finished: widescreen from the start, fitted pictures, styled tables,
native charts, notes, and slide duplicate/delete/move (which python-pptx does not offer).

Use from a script:  sys.path.insert(0, "<this folder>"); from helpers import *
"""
import copy
import math

from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION
from pptx.util import Emu, Inches, Pt

__all__ = ["new_deck", "set_title", "add_bullets", "add_textbox", "estimate_font_size", "add_picture_fit", "add_table",
           "add_bar_chart", "set_notes", "duplicate_slide", "delete_slide", "move_slide"]

WIDE_W, WIDE_H = Inches(13.333), Inches(7.5)


def new_deck(template=None, widescreen=True):
    """Presentation from a template (or the default one) in 16:9.

    The default template is 4:3; resizing the slide alone leaves every layout placeholder where it was, so the
    horizontal geometry of the master and the layouts is scaled with the page.
    """
    prs = Presentation(template) if template else Presentation()
    if widescreen and prs.slide_width != WIDE_W:
        factor = WIDE_W / prs.slide_width
        prs.slide_width, prs.slide_height = WIDE_W, WIDE_H
        for container in [prs.slide_master, *prs.slide_layouts]:
            for shape in container.shapes:
                # only shapes with their own geometry: the others inherit it from the master, already scaled
                if shape._element.xpath("./p:spPr/a:xfrm"):
                    shape.left, shape.width = int(shape.left * factor), int(shape.width * factor)
    return prs


def set_title(slide, text):
    if slide.shapes.title is None:
        raise ValueError("this layout has no title placeholder: pick another layout")
    slide.shapes.title.text = text


def add_bullets(slide, items, placeholder_idx=1, size_pt=None):
    """Fill the body placeholder. `items` are strings, or (text, level) pairs for indented sub-points."""
    frame = slide.placeholders[placeholder_idx].text_frame
    frame.clear()
    for i, item in enumerate(items):
        text, level = (item, 0) if isinstance(item, str) else item
        paragraph = frame.paragraphs[0] if i == 0 else frame.add_paragraph()
        paragraph.text = text
        paragraph.level = level
        if size_pt:
            for run in paragraph.runs:
                run.font.size = Pt(size_pt)
    return frame


def estimate_font_size(text, width_in, height_in, max_pt=28, min_pt=12, line_spacing=1.2):
    """Largest size in [min_pt, max_pt] at which `text` should fit the box (average glyph width = 0.5 em).

    A heuristic for proportional fonts; confirm with render.py. Returns min_pt when even that is too big, which is
    the signal to shorten the text or split the slide.
    """
    for size in range(int(max_pt), int(min_pt) - 1, -1):
        chars_per_line = max(1, int(width_in * 72 / (size * 0.5)))
        lines = sum(max(1, math.ceil(len(part) / chars_per_line)) for part in text.split("\n"))
        if lines * size * line_spacing <= height_in * 72:
            return size
    return min_pt


def add_textbox(slide, text, left_in, top_in, width_in, height_in, size_pt=None, bold=False, color=None, align=None):
    """Word-wrapped text box; with no `size_pt` the size is estimated to fit the box."""
    box = slide.shapes.add_textbox(Inches(left_in), Inches(top_in), Inches(width_in), Inches(height_in))
    frame = box.text_frame
    frame.word_wrap = True
    size = size_pt or estimate_font_size(text, width_in, height_in)
    for i, line in enumerate(text.split("\n")):
        paragraph = frame.paragraphs[0] if i == 0 else frame.add_paragraph()
        run = paragraph.add_run()
        run.text = line
        run.font.size = Pt(size)
        run.font.bold = bold
        if color:
            run.font.color.rgb = RGBColor.from_string(color)
        if align is not None:
            paragraph.alignment = align
    return box


def add_picture_fit(slide, path, left_in, top_in, width_in, height_in, alt):
    """Picture scaled to fit inside the box without distortion, centered in it, with alt text."""
    from PIL import Image
    with Image.open(path) as img:
        ratio = img.width / img.height
    w, h = (width_in, width_in / ratio) if width_in / height_in < ratio else (height_in * ratio, height_in)
    pic = slide.shapes.add_picture(str(path), Inches(left_in + (width_in - w) / 2), Inches(top_in + (height_in - h) / 2),
                                   Inches(w), Inches(h))
    pic._element.nvPicPr.cNvPr.set("descr", alt)
    return pic


def add_table(slide, header, rows, left_in, top_in, width_in, row_height_in=0.4, size_pt=14, header_fill="1F4E78",
              header_color="FFFFFF", align_right=()):
    """Table with a coloured header row; `align_right` lists numeric column indexes."""
    from pptx.enum.text import PP_ALIGN
    shape = slide.shapes.add_table(len(rows) + 1, len(header), Inches(left_in), Inches(top_in), Inches(width_in),
                                   Inches(row_height_in * (len(rows) + 1)))
    table = shape.table
    for r, values in enumerate([header, *rows]):
        for c, value in enumerate(values):
            cell = table.cell(r, c)
            cell.text = str(value)
            for paragraph in cell.text_frame.paragraphs:
                if c in align_right:
                    paragraph.alignment = PP_ALIGN.RIGHT
                for run in paragraph.runs:
                    run.font.size = Pt(size_pt)
                    run.font.bold = r == 0
                    if r == 0:
                        run.font.color.rgb = RGBColor.from_string(header_color)
            if r == 0:
                cell.fill.solid()
                cell.fill.fore_color.rgb = RGBColor.from_string(header_fill)
    return shape


def add_bar_chart(slide, categories, series, left_in, top_in, width_in, height_in, title=None, kind=XL_CHART_TYPE.COLUMN_CLUSTERED):
    """Native, editable chart. `series` maps a series name to its list of values (same length as `categories`)."""
    data = CategoryChartData()
    data.categories = categories
    for name, values in series.items():
        data.add_series(name, values)
    chart = slide.shapes.add_chart(kind, Inches(left_in), Inches(top_in), Inches(width_in), Inches(height_in), data).chart
    chart.has_legend = len(series) > 1
    if chart.has_legend:
        chart.legend.position = XL_LEGEND_POSITION.BOTTOM
        chart.legend.include_in_layout = False
    if title:
        chart.has_title = True
        chart.chart_title.text_frame.text = title
    return chart


def set_notes(slide, text):
    """Speaker notes."""
    slide.notes_slide.notes_text_frame.text = text


def duplicate_slide(prs, index):
    """Append a copy of slide `index` (same layout, shapes, pictures, links). A copied chart shares its data part with
    the original: re-create the chart instead of editing both."""
    src = prs.slides[index]
    dst = prs.slides.add_slide(src.slide_layout)
    for shape in list(dst.shapes):
        shape._element.getparent().remove(shape._element)
    mapping = {}
    for rel in src.part.rels.values():
        if rel.reltype.endswith("/notesSlide") or rel.reltype.endswith("/slideLayout"):
            continue
        if rel.is_external:
            mapping[rel.rId] = dst.part.rels.get_or_add_ext_rel(rel.reltype, rel.target_ref)
        else:
            mapping[rel.rId] = dst.part.rels.get_or_add(rel.reltype, rel.target_part)
    for shape in src.shapes:
        element = copy.deepcopy(shape._element)
        for node in element.iter():
            for attr, value in list(node.attrib.items()):
                if attr.split("}")[-1] in ("embed", "link", "id", "pict") and attr.startswith("{http://schemas.openxmlformats.org/officeDocument/2006/relationships}"):
                    if value in mapping:
                        node.set(attr, mapping[value])
        dst.shapes._spTree.insert_element_before(element, "p:extLst")
    if src.has_notes_slide:
        dst.notes_slide.notes_text_frame.text = src.notes_slide.notes_text_frame.text
    return dst


def delete_slide(prs, index):
    """Remove slide `index` from the deck and drop its relationship."""
    ids = prs.slides._sldIdLst
    slide_id = list(ids)[index]
    prs.part.drop_rel(slide_id.rId)
    ids.remove(slide_id)


def move_slide(prs, old_index, new_index):
    """Reorder: the slide at `old_index` ends up at `new_index`."""
    ids = prs.slides._sldIdLst
    items = list(ids)
    element = items[old_index]
    ids.remove(element)
    ids.insert(new_index, element)
